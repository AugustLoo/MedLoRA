"""第 6 周周进展汇报 PPT (课题组统一模板 Weekly Report_Templet): docs/weekly/week6/Topic6_WeeklyReport6.pptx
只往模板原有的文本框 / 表格里填字 (保留字体、字号、颜色), 每格「中文一行 + 英文一行」; 数字从 results/ 的快照读。
用法: python scripts/make_week6_ppt.py <模板.pptx>
"""
import copy
import json
import sys
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Inches

REPO = Path(__file__).resolve().parents[1]
R = REPO / "results"
TEMPLATE = Path(sys.argv[1])
OUT = REPO / "docs/weekly/week6/Topic6_WeeklyReport6.pptx"

EV = json.loads((R / "interns2_3stage_eval_2026-10-19.json").read_text(encoding="utf-8"))["models"]
CI_A = json.loads((R / "interns2_ci_3stage_vsA_2026-10-19.json").read_text(encoding="utf-8"))["tests"]
CI_B = json.loads((R / "interns2_ci_3stage_CvsB_2026-10-19.json").read_text(encoding="utf-8"))["tests"]
TR = json.loads((R / "interns2_3stage_train_2026-10-18.json").read_text(encoding="utf-8"))
BASE, SIX, M1, MA, MB, MC = (EV[k] for k in ("原版", "S2-6datasets", "只做①", "A (②)", "B (①→②)", "C (①→②→③)"))
D1, D2, D3 = TR["data"]["stage1"], TR["data"]["stage2_zh_dialog"], TR["data"]["stage3_reasoning"]
TOK_M = D1["_total"]["tokens"] / 1e6
N_ZH = D2["huatuo_sft"]["kept"] + D2["disc_med_sft"]["kept"]
N_R = sum(D3[k]["kept"] for k in ("cmexam", "medmcqa", "medical_o1_zh", "medical_o1_en"))


def r1(x):
    return str(Decimal(str(x)).quantize(Decimal("0.1"), ROUND_HALF_UP))


def sg(x):
    s = r1(x)
    return s if s.startswith("-") else "+" + s


def d(ci, test, tag):
    x = ci[test][tag]
    return f"{sg(x['diff_vs_ref'])} [{sg(x['diff_ci95'][0])}, {sg(x['diff_ci95'][1])}]"


BA_SL = d(CI_A, "SLAKE closed", "interns2_s2b")
BA_PATH = d(CI_A, "PathVQA closed", "interns2_s2b")
CB_PNEU = d(CI_B, "PneumoniaMNIST closed", "interns2_s3c")
N_BA_NOISE = sum(not CI_A[k]["interns2_s2b"]["significant"] for k in (
    "SLAKE closed", "SLAKE open", "SLAKE 中文 open", "PathVQA closed", "MedQA acc", "PneumoniaMNIST closed"))


# ---------------------------------------------------------------- 填字工具 (保留原有段落 / 字的格式)
def fill(tf, lines):
    """把文本框 / 单元格的段落换成 lines (每项一段); 段落不够就复制最后一个非空段落, 多了就删掉。"""
    paras = list(tf.paragraphs)
    proto = next((p for p in reversed(paras) if p.runs), paras[0])
    for p in paras[len(lines):]:
        p._p.getparent().remove(p._p)
    paras = paras[:len(lines)]
    while len(paras) < len(lines):
        el = copy.deepcopy(proto._p)
        paras[-1]._p.addnext(el)
        paras = list(tf.paragraphs)
    for p, text in zip(tf.paragraphs, lines):
        runs = p.runs
        if not runs:  # 空段落: 借原型段落的第一个字的格式
            if text == "":
                continue
            p._p.append(copy.deepcopy(proto.runs[0]._r))
            runs = p.runs
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)


def shape(slide, name):
    return next(s for s in slide.shapes if s.name == name)


def table(slide):
    return next(s for s in slide.shapes if s.has_table).table


def set_rows(tbl, n_body, row_h=None):
    """表格正文改成 n_body 行 (复制最后一行或删掉多余的), 可统一行高。"""
    rows = tbl._tbl.tr_lst
    while len(rows) - 1 > n_body:
        rows[-1].getparent().remove(rows[-1])
        rows = tbl._tbl.tr_lst
    while len(rows) - 1 < n_body:
        rows[-1].addnext(copy.deepcopy(rows[-1]))
        rows = tbl._tbl.tr_lst
    if row_h:
        for r in tbl.rows:
            r.height = Emu(int(Inches(row_h)))


def fill_table(tbl, body):
    for i, row in enumerate(body, start=1):
        for j, (zh, en) in enumerate(row):
            fill(tbl.cell(i, j).text_frame, [zh, en])


def boxes(slide, items, takeaway):
    for i, (body_lines) in enumerate(items):
        fill(shape(slide, f"body-{i}").text_frame, body_lines)
    fill(shape(slide, "takeaway").text_frame, takeaway)


prs = Presentation(str(TEMPLATE))
S = prs.slides
NOTES = []

# ---------------------------------------------------------------- 1 封面
fill(shape(S[0], "details").text_frame, [
    "医学多模态模型的参数高效增量预训练与指令微调",
    "Parameter-Efficient Incremental Pretraining and Instructional Fine-Tuning of Medical Multimodal Models",
    "Chunqian Loo    导师 / Supervisor：徐蕾",
    "第 6 周 / Week 6    2026-10-22",
    "",
    "本周核心进展：三步训练的四个模型全部评完；分数来自微调，增量预训练没有可测提升",
    "Main update: all four models evaluated; gains come from fine-tuning, pretraining adds nothing measurable",
])
NOTES.append("这周按苏学长的方案，在 35B 上做了三步训练：先读医学教材（预训练），再练题（微调），最后练带解析的考题（推理训练）。"
             "四个模型都训完、评完了。一句话结论：分数几乎都来自微调；加预训练、加推理训练都没有可测的提升。")

# ---------------------------------------------------------------- 2 本周进展
t = table(S[1])
fill_table(t, [
    [("三步训练（苏学长的方案）", "Three-stage training (Su's plan)"),
     ("完成：4 个模型训完、评完", "Done: 4 models evaluated"),
     ("结果快照与置信区间已推 GitHub", "Results and CIs pushed to GitHub")],
    [("中文数据：先测再训", "Chinese data: measure, then train"),
     ("部分完成", "Partly done"),
     ("基线已测；SLAKE 中文题未加", "Baseline done; SLAKE zh items next")],
    [("种子 43 重训 A、B", "Seed-43 retrain of A and B"),
     ("进行中，已暂停", "In progress, paused"),
     ("A 训到一半（断点已存），卡让给世龙", "A half done (checkpoint); GPUs back to Shilong")],
])
fill(shape(S[1], "takeaway").text_frame, [
    "主要阻碍：这几天卡要跑世龙的服务，训练暂停    所需支持：约 18 小时显卡、临时硬盘",
    "Main blocker: GPUs serve Shilong's model this week; training paused    Support needed: ~18 GPU hours, disk space",
])
NOTES.append("上周说的三件事：三步训练已经全部完成；中文数据部分完成，基线测了、中文对话加进了第 2 步，但 SLAKE 中文训练题还没加；"
             "种子 43 重训已经开始，A 训到一半、断点已存，但这几天卡要跑世龙的服务，所以先暂停。需要的支持：卡空出来后约 18 小时显卡，以及硬盘临时空间。")

# ---------------------------------------------------------------- 3 核心挑战与研究问题
boxes(S[2], [
    ["输入医学图像 + 中英文问题，输出短答或描述；35B，LoRA 只训 0.054%",
     "Medical image + EN/ZH question → short answer or description; 35B, LoRA 0.054%"],
    ["微调后胸片描述会漏报、和问答自相矛盾；3B 上预训练没有增益",
     "After fine-tuning, X-ray descriptions miss findings; CPT gave nothing at 3B"],
    ["RQ1：微调前加增量预训练（读教材），能否提升医学看图问答？",
     "RQ1: does incremental pretraining before fine-tuning improve medical VQA?",
     "本周验证：B（预训练 + 微调）对比 A（只微调） / This week's test: B vs A"],
], [
    "研究假设：相同 LoRA 设置、种子 42 下，预训练 + 微调比只微调的 SLAKE / PathVQA / MedQA 更高",
    "Hypothesis: with the same LoRA and seed 42, CPT + fine-tuning beats fine-tuning on SLAKE / PathVQA / MedQA",
])
NOTES.append("任务：给医学图像和问题，模型回答或描述。核心挑战是可靠性：胸片描述会漏报、会自相矛盾；另外 3B 上预训练没用。"
             "这周的研究问题是：在微调之前加一步读教材的预训练，看图问答会不会更好？用 B 对比 A 来验证。")

# ---------------------------------------------------------------- 4 文献
t = table(S[3])
fill_table(t, [
    [("LoRA Learns Less…（2024）", "Biderman et al., TMLR"),
     ("LoRA 预训练学得少、忘得少", "LoRA CPT learns less, forgets less"),
     ("没在医学多模态上验证", "Not tested on medical VLMs")],
    [("Lost in the Hype（2026）", "Zhu et al., arXiv 2604.08333"),
     ("训视觉塔只差约 ±0.5 分", "Vision tower: about ±0.5 points"),
     ("「看得到、说不出」，没给修法", "Sees but does not say; no fix")],
    [("Consistent but Dangerous", "CVPR 2026 Workshop"),
     ("可能不看图，建议去图对照", "May ignore images; no-image check"),
     ("没有让问答与描述一致的方法", "No fix for answer–description gaps")],
])
fill(shape(S[3], "takeaway").text_frame, [
    "研究空白：MoE 多模态模型上预训练有没有用、问答与描述是否一致    本周阅读：PU-DPO（报告漏写 → 少报）",
    "Gap: does CPT help MoE VLMs; answer–description consistency    Reading: PU-DPO (models learn omissions)",
])
NOTES.append("三篇文献：LoRA 做预训练学得少、忘得少，和我们两个规模都没提升一致；训练视觉编码器只差约 0.5 分，支持冻结；"
             "微调后模型可能不看图，所以我们加了去掉图像的对照。还读了 PU-DPO：医生报告会漏写，模型学会少报。数字汇报前要再核对原文。")

# ---------------------------------------------------------------- 5 拟议方法
boxes(S[4], [
    ["[图像 + 问题] → 教材预训练 → 问答微调 → 带解析考题 → [回答 + 可靠性评测]",
     "[Image + question] → textbook CPT → Q&A fine-tuning → explained exams → [answer]"],
    ["最接近 Lingshu 的多阶段全参数训练；我们只用一个 0.054% 补丁，每步加一类数据",
     "Closest: Lingshu's multi-stage full training; ours: one 0.054% adapter, one data type per stage"],
    ["C1：35B MoE 上预训练 / 推理训练的受控消融；C2：可靠性评测（去图、漏报）",
     "C1: controlled CPT / reasoning ablation on a 35B MoE; C2: reliability evaluation",
     "验证：B − A、C − B 配对置信区间 / Evidence: paired CIs"],
], [
    "本周方法更新：一次微调改成三步训练；评测加上「去掉图像」对照和置信区间",
    "Method update: one fine-tuning run became three stages; added a no-image control and confidence intervals",
])
NOTES.append("方法：同一个 LoRA 补丁，三步接着训，每步只多加一类数据，这样能看出每一步的作用。贡献一是 35B 和 3B 上预训练的受控对比，"
             "贡献二是可靠性评测：去掉图像、漏报、问答和描述矛盾。")

# ---------------------------------------------------------------- 6 实验计划
t = table(S[5])
fill_table(t, [
    [("问题与对比", "Question / comparison"),
     ("A（只微调）对比 B（预训练 + 微调）；C（+ 推理训练）对比 B",
      "A (fine-tune) vs B (CPT + fine-tune); C (+ reasoning) vs B")],
    [("数据与划分", "Data / split"),
     (f"教材 {TOK_M:,.1f}M token；问答 15,953 + 中文对话 {N_ZH:,}；考题 {N_R:,}",
      f"{TOK_M:,.1f}M tokens; 15,953 + {N_ZH:,} Q&A; {N_R:,} exam items; test sets isolated")],
    [("指标与控制变量", "Metrics / controls"),
     ("看图 / 考题准确率（↑）、胸片漏报（↓）；LoRA、种子 42 固定",
      "VQA / exam accuracy (↑), X-ray omission (↓); LoRA r16, seed 42 fixed")],
    [("消融与成功标准", "Ablation / criterion"),
     ("去掉第 1 步（A）、第 3 步（B）；区间不跨 0 才算有差别",
      "Remove stage 1 (A) or 3 (B); success = 95% CI excludes 0")],
])
fill(shape(S[5], "takeaway").text_frame, [
    "本周实验：种子 43 重训 A 已到一半，暂停中；B 未开始    复现记录：脚本与结果在 GitHub MedLoRA",
    "This week's run: seed-43 A half done, paused; B not started    Reproducibility: scripts and results on GitHub",
])
NOTES.append("实验设计：A 对比 B 看预训练，C 对比 B 看推理训练。所有设置固定，只差一步。成功标准事先定好：配对 95% 置信区间不跨 0 才算有差别。")

# ---------------------------------------------------------------- 7 实验结果 (多加一行: 只做预训练)
t = table(S[6])
set_rows(t, 4, row_h=0.72)
fill_table(t, [
    [("A 只做微调（基线）", "A: fine-tune only (baseline)"),
     (f"{r1(MA['slake_closed'])} / {r1(MA['slake_open_em'])}", "SLAKE closed / open"),
     (f"{r1(MA['cxr_desc_miss_pct'])}%", "X-ray omission"),
     ("interns2_s2a", "5,241 steps")],
    [("B 预训练 + 微调（拟议）", "B: CPT + fine-tune (proposed)"),
     (f"{r1(MB['slake_closed'])} / {r1(MB['slake_open_em'])}", f"B − A {BA_SL}"),
     (f"{r1(MB['cxr_desc_miss_pct'])}%", "X-ray omission"),
     ("interns2_s2b", "from interns2_cpt1")],
    [("C 三步都做（加推理训练）", "C: + reasoning training"),
     (f"{r1(MC['slake_closed'])} / {r1(MC['slake_open_em'])}", "C − B within noise"),
     (f"{r1(MC['cxr_desc_miss_pct'])}%", "X-ray omission"),
     ("interns2_s3c", "812 steps")],
    [("只做预训练（不微调）", "CPT only, no fine-tuning"),
     (f"{r1(M1['slake_closed'])} / {r1(M1['slake_open_em'])}", "near the base model"),
     (f"{r1(M1['cxr_desc_miss_pct'])}%", "X-ray omission"),
     ("interns2_cpt1", "1,502 steps")],
])
fill(shape(S[6], "takeaway").text_frame, [
    f"观察：B − A 六项中 {N_BA_NOISE} 项在噪声内，C ≈ B    解释：分数来自微调    局限：单种子，胸片待种子 43",
    f"Observation: B − A within noise on {N_BA_NOISE} of 6, C ≈ B    Interpretation: gains from fine-tuning    Limitation: one seed",
])
NOTES.append(f"结果：A、B、C 三个微调过的模型几乎一样高，SLAKE 是非题 93.5 到 94.2。B 减 A 在 SLAKE 是非题上是 {BA_SL}，区间跨 0，在噪声内；"
             f"六项里只有病理图 {BA_PATH} 有差别。C 减 B 只有儿童肺炎 {CB_PNEU} 刚好过线。只做预训练不微调，看图题在原版附近。"
             f"去掉图像后 B 的是非题从 {r1(MB['slake_closed'])} 掉到 {r1(MB['slake_noimg_closed'])}，说明模型真的在看图。"
             f"成人胸片描述漏报 A/B/C 是 {r1(MA['cxr_desc_miss_pct'])}%、{r1(MB['cxr_desc_miss_pct'])}%、{r1(MC['cxr_desc_miss_pct'])}%，"
             f"比上周 S2-6datasets 的 {r1(SIX['cxr_desc_miss_pct'])}% 差，可能是中文对话也可能是种子，要用种子 43 确认。")

# ---------------------------------------------------------------- 8 预期成果与阶段结论
boxes(S[7], [
    ["论文实验章节：3B、35B 预训练消融和可靠性评测；同意的话加原型系统",
     "Thesis chapters: CPT ablation at 3B and 35B, reliability tests; a prototype if approved"],
    [f"文字预训练、推理训练都没有可测提升；去掉图像后是非题 {r1(MB['slake_noimg_closed'])}，模型在看图",
     "No measurable gain from CPT or reasoning; without the image the models drop to chance"],
    ["胸片描述变差是中文对话还是随机种子？→ 种子 43 重训 A、B",
     "Is the X-ray drop from the dialogues or the seed? → seed-43 retrain of A and B"],
], [
    "最终交付：代码（GitHub MedLoRA）、LoRA 补丁、数据卡 / 模型卡、演示系统（待确认）、论文",
    "Final deliverables: code, LoRA adapters, data / model cards, demo (to confirm), thesis",
])
NOTES.append("当前结论：在我们这个设置下，文字预训练和推理训练都没有可测提升，两个规模结果一致，这本身可以写成论文的结论。"
             "还不确定的是胸片描述变差的原因，以及带图的预训练会不会有用。")

# ---------------------------------------------------------------- 9 时间安排
t = table(S[8])
fill_table(t, [
    [("第1–4周", "W1–4"), ("数据与基线：3B、35B 基座", "Data and baselines: 3B, 35B"),
     ("results/ 快照、周报 1–3", "Result snapshots, reports 1–3"), ("完成", "Done")],
    [("第5–8周", "W5–8"), ("35B 微调、三步训练", "35B fine-tuning, 3 stages"),
     ("评测结果、置信区间", "Evaluation results, CIs"), ("进行中", "In progress")],
    [("第9–12周", "W9–12"), ("种子 43、CheXbert、中文看图", "Seed 43, CheXbert, Chinese VQA"),
     ("图表与错误分析", "Figures, error analysis"), ("未开始", "Not started")],
    [("第13–16周", "W13–16"), ("论文与演示", "Thesis and demo"),
     ("初稿、原型系统", "Draft, prototype"), ("未开始", "Not started")],
])
fill(shape(S[8], "takeaway").text_frame, [
    "下周承诺：卡空出来后接着训 A、再训 B（共约 18 小时）    风险预案：排不上就先写实验章节",
    "Next week: resume A, then B, when GPUs free up (~18 h)    Fallback: write the experiment chapter first",
])
NOTES.append("时间安排：前四周做完数据和基线，第 5 到 8 周是核心方法，现在三步训练已经评完。种子 43 的 A 已经训到一半，卡空出来后接着训，再训 B；排不上就先写论文的实验章节。")

# ---------------------------------------------------------------- 10 讨论与决策
boxes(S[9], [
    ["要不要做带图的预训练？A：做（约 2–3 晚显卡）；B：不做，写成结论。建议先做种子 43 再定",
     "Image-text CPT? A: do it (~2–3 GPU nights); B: report the finding. Suggest seed 43 first"],
    ["显卡约 18 小时（世龙服务空出时）；临时硬盘；论文能否加原型系统一章",
     "~18 GPU hours when Shilong's service allows; temporary disk; a prototype chapter?"],
    ["会议决策：（会后填写）", "Decision: (to fill after the meeting)", "执行事项：（会后填写） / Action: (after the meeting)"],
], [
    "每周材料：汇报 PPT、实验日志（results/README.md）、代码与配置（GitHub AugustLoo/MedLoRA）",
    "Weekly package: slides, experiment log (results/README.md), code and config (GitHub AugustLoo/MedLoRA)",
])
NOTES.append("想请老师决定：要不要做带图的预训练。我的建议是先用种子 43 确认现在的结论，再决定。需要的支持：约 18 小时显卡、硬盘临时空间，以及论文能否加原型系统那一章。")

for slide, note in zip(S, NOTES):
    slide.notes_slide.notes_text_frame.text = note

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(str(OUT))
print("written", OUT)
